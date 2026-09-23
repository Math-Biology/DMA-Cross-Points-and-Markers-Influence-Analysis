# LINKED-TO: [REQ-CPM-IA-P26.0002]
"""
data_ingestion.py — Autonomous Raw-Variation Ingestion for the CPM-IA component.

Reads input data from the configured source (CSV, Excel, or database), validates
its schema, enforces data quality rules, and returns a clean DataFrame.
All field validation occurs here before any downstream computation (G-03).
"""

from __future__ import annotations

import logging
import re
from abc import ABC, abstractmethod
from dataclasses import dataclass
from pathlib import Path
from typing import Tuple

import pandas as pd

from src.config_loader import (
    CpmIaConfig,
    CsvInputConfig,
    DbInputConfig,
    ExcelInputConfig,
)

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Column name constants — confirmed contract with upstream Raw Variation Computation
# Confirmed by [IR] 2026-08-27 (updated 2026-08-28): visit_id, marker, point, percentage_variation, visit_date
# ---------------------------------------------------------------------------

COL_VISIT = "visit_id"
COL_MARKER = "marker"
COL_POINT = "point"
COL_RAW_VARIATION = "percentage_variation"
COL_VISIT_DATE = "visit_date"  # kept as string/object; no datetime parsing at ingestion

REQUIRED_COLUMNS: frozenset[str] = frozenset({
    COL_VISIT,
    COL_MARKER,
    COL_POINT,
    COL_RAW_VARIATION,
    COL_VISIT_DATE,
})

_KEY_ID_COLUMNS: tuple[str, ...] = (COL_VISIT, COL_MARKER, COL_POINT)

# Valid anatomical point patterns:
#   new format (4 components): LH - 5 - e - Small intestine
#   short format (3 components): LH - 3 - c  (4th component optional)
#   legacy format: SomeName [3]
_VALID_POINT_PATTERN = re.compile(
    r"^[LR][HF]\s*-\s*\d+\s*-\s*\S+(\s*-\s*.+)?|.+\s*\[\d+\]\s*$"
)


class IngestionError(Exception):
    """Raised when the input is missing, malformed, or fails schema validation."""


# ---------------------------------------------------------------------------
# Data quality report
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class DataQualityReport:
    total_rows: int
    rows_dropped_empty_keys: int
    invalid_point_names: tuple   # Tuple[str, ...] of invalid point names found
    rows_affected_invalid_points: int   # rows dropped or warned
    drop_mode: bool              # True when rows were dropped


# ---------------------------------------------------------------------------
# DataSource abstraction
# ---------------------------------------------------------------------------

class DataSource(ABC):
    @abstractmethod
    def load(self) -> pd.DataFrame:
        """Return a raw long-format DataFrame (before validation)."""
        ...


class CsvSource(DataSource):
    def __init__(self, config: CsvInputConfig) -> None:
        self._config = config

    def load(self) -> pd.DataFrame:
        path = Path(self._config.path)
        if not path.exists():
            raise IngestionError(f"Input file not found: {path}")
        if path.suffix.lower() != ".csv":
            raise IngestionError(
                f"Unsupported input format: {path.suffix!r}; expected .csv"
            )
        return pd.read_csv(path, encoding="utf-8")


class ExcelSource(DataSource):
    def __init__(self, config: ExcelInputConfig) -> None:
        self._config = config

    def load(self) -> pd.DataFrame:
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            raise IngestionError(
                "openpyxl is required for Excel ingestion: pip install openpyxl"
            )

        frames = []
        for file_cfg in self._config.files:
            path = Path(file_cfg.path)
            if not path.exists():
                raise IngestionError(f"Excel input file not found: {path}")

            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active

            # Read header row
            headers = [cell.value for cell in ws[1]]

            # First column = marker; remaining columns = anatomical points (skip None/empty headers)
            point_cols = []
            for col_idx, h in enumerate(headers[1:], start=1):
                if h is not None and str(h).strip():
                    point_cols.append((col_idx, str(h).strip()))

            # Build rows
            rows = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                marker_val = row[0]
                if marker_val is None:
                    continue
                marker_str = str(marker_val).strip()
                if not marker_str:
                    continue
                for col_idx, point_name in point_cols:
                    val = row[col_idx] if col_idx < len(row) else None
                    rows.append({
                        COL_VISIT: file_cfg.visit_id,
                        COL_MARKER: marker_str,
                        COL_POINT: point_name,
                        COL_RAW_VARIATION: val,
                        COL_VISIT_DATE: file_cfg.visit_date,
                    })

            if rows:
                frames.append(pd.DataFrame(rows))

        if not frames:
            # Return empty DataFrame with correct columns
            return pd.DataFrame(columns=list(REQUIRED_COLUMNS))

        return pd.concat(frames, ignore_index=True)


class PostgresSource(DataSource):
    def __init__(self, config: DbInputConfig) -> None:
        self._config = config

    def load(self) -> pd.DataFrame:
        raise IngestionError(
            f"Postgres source not configured: dsn_env={self._config.dsn_env!r}. "
            "Install psycopg2 and set the DSN environment variable before use."
        )


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def load_data(config: CpmIaConfig) -> Tuple[pd.DataFrame, DataQualityReport]:
    """
    Read and validate the raw-variation data from the configured source.

    Returns:
        A tuple (df_clean, DataQualityReport).
        df_clean is a validated long-format DataFrame with COL_RAW_VARIATION as float64
        and COL_VISIT_DATE as string/object.

    Raises:
        IngestionError on any failure.
    """
    # Step 1 — dispatch to appropriate DataSource
    input_type = config.input.type
    if input_type == "csv":
        if config.input.csv is None:
            raise IngestionError("Input type is 'csv' but no <csv> config block found")
        source: DataSource = CsvSource(config.input.csv)
    elif input_type == "excel":
        if config.input.excel is None:
            raise IngestionError("Input type is 'excel' but no <excel> config block found")
        source = ExcelSource(config.input.excel)
    elif input_type == "db":
        if config.input.db is None:
            raise IngestionError("Input type is 'db' but no <db> config block found")
        source = PostgresSource(config.input.db)
    else:
        raise IngestionError(f"Unknown input type: {input_type!r}")

    df = source.load()

    # Step 2 — required column presence (G-03)
    missing = REQUIRED_COLUMNS - set(df.columns)
    if missing:
        raise IngestionError(f"Missing required column(s): {sorted(missing)}")

    # Step 3 — numeric cast of COL_RAW_VARIATION
    # Genuine NaN (empty cell — missing measurement) is allowed and passes through.
    # Only rows where the original value is non-NaN but numeric coercion fails are rejected.
    coerced = pd.to_numeric(df[COL_RAW_VARIATION], errors="coerce")
    invalid_mask = coerced.isna() & df[COL_RAW_VARIATION].notna()
    if invalid_mask.any():
        bad_indices = df.index[invalid_mask].tolist()
        raise IngestionError(
            f"Non-numeric values in {COL_RAW_VARIATION!r} at rows: {bad_indices}"
        )
    df[COL_RAW_VARIATION] = coerced.astype("float64")

    # Step 4 — NaN check in key identifier columns
    bad_mask = df[list(_KEY_ID_COLUMNS)].isna().any(axis=1)
    if bad_mask.any():
        raise IngestionError(
            f"NaN in key columns at rows: {df.index[bad_mask].tolist()}"
        )

    total_rows = len(df)

    # Step 5 — drop rows with whitespace-only key fields
    for col in _KEY_ID_COLUMNS:
        df[col] = df[col].astype(str)
    whitespace_mask = (
        df[COL_VISIT].str.strip().eq("") |
        df[COL_MARKER].str.strip().eq("") |
        df[COL_POINT].str.strip().eq("")
    )
    rows_dropped_empty = int(whitespace_mask.sum())
    if rows_dropped_empty:
        logger.warning(
            "Dropping %d rows with whitespace-only key fields", rows_dropped_empty
        )
        df = df[~whitespace_mask].copy()

    # Step 6 — anatomical point validation
    unique_points = df[COL_POINT].unique()
    invalid_points = [
        p for p in unique_points
        if not _VALID_POINT_PATTERN.match(str(p).strip())
    ]
    invalid_point_set = set(invalid_points)
    invalid_rows_mask = df[COL_POINT].isin(invalid_point_set)
    rows_affected = int(invalid_rows_mask.sum())

    if invalid_points:
        if config.data_quality.drop_invalid_points:
            logger.warning(
                "Dropping %d rows with invalid anatomical point names: %s",
                rows_affected,
                sorted(invalid_points),
            )
            df = df[~invalid_rows_mask].copy()
        else:
            logger.warning(
                "Found %d rows with invalid anatomical point names (kept): %s",
                rows_affected,
                sorted(invalid_points),
            )

    dq_report = DataQualityReport(
        total_rows=total_rows,
        rows_dropped_empty_keys=rows_dropped_empty,
        invalid_point_names=tuple(sorted(invalid_points)),
        rows_affected_invalid_points=rows_affected,
        drop_mode=config.data_quality.drop_invalid_points,
    )

    logger.info(
        "CPM-IA input loaded | rows=%d | columns=%s | visits=%d | markers=%d | points=%d",
        len(df),
        sorted(df.columns.tolist()),
        df[COL_VISIT].nunique(),
        df[COL_MARKER].nunique(),
        df[COL_POINT].nunique(),
    )

    return df, dq_report
